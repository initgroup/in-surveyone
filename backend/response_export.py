from __future__ import annotations

import io
import json
import re
from collections.abc import Callable
from dataclasses import dataclass
from datetime import timezone
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.errors import DomainValidationError
from backend.models import Question, ResponseWorkbookRequest, SurveyDetail, SurveyResponseView
from backend.memory_store import UserRecord
from backend.stores import SurveyStore


@dataclass(frozen=True)
class ExportColumn:
    label: str
    question_id: str
    section: str
    original_title: str
    value: Callable[[dict[str, Any]], Any]


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def _label(items: list, value: Any) -> str:
    token = str(value) if value is not None else ""
    match = next((item for item in items if str(item.id) == token), None)
    if match is not None:
        return match.label
    if token.isdecimal() and int(token) < len(items):
        return items[int(token)].label
    return token


def _display(question: Question, value: Any) -> Any:
    if value is None or value == "" or value == [] or value == {}:
        return ""
    kind = question.type.value
    if kind in {"single", "single_choice", "dropdown"}:
        return _label(question.options, value)
    if kind in {"multiple", "multiple_choice"}:
        return "\n".join(_label(question.options, item) for item in (value if isinstance(value, list) else [value]))
    if kind in {"matrix", "matrix_single", "matrix_multiple"} and isinstance(value, dict):
        return "\n".join(
            f"{_label(question.settings.rows, row)}: " + ", ".join(
                _label(question.settings.columns, item) for item in (selected if isinstance(selected, list) else [selected])
            ) for row, selected in value.items()
        )
    if isinstance(value, list) and kind in {"table", "repeat_table"}:
        return "\n".join(
            f"{question.settings.rows[index].label if index < len(question.settings.rows) else str(index + 1) + '행'}: " + " · ".join(
                f"{question.columns[column] if column < len(question.columns) else str(column + 1) + '열'}={cell}"
                for column, cell in enumerate(row if isinstance(row, list) else list(row.values()) if isinstance(row, dict) else [row])
            ) for index, row in enumerate(value)
        )
    return _json(value) if isinstance(value, (list, dict)) else value


def _columns(survey: SurveyDetail, request: ResponseWorkbookRequest) -> list[ExportColumn]:
    known = {str(q.id) for section in survey.sections for q in section.questions if q.type.value not in {"notice", "text"}}
    if request.question_ids is not None and not set(request.question_ids).issubset(known):
        raise DomainValidationError("표시 문항이 변경되었습니다. 응답 탐색을 새로고침해 주세요.")
    selected = set(request.question_ids) if request.mode == "visible" and request.question_ids is not None else known
    columns: list[ExportColumn] = []
    number = 0
    for section in survey.sections:
        for q in section.questions:
            if q.type.value in {"notice", "text"}:
                continue
            numbered = survey.settings.show_question_numbers and not q.settings.exclude_from_numbering
            if numbered:
                number += 1
            title = re.sub(r"^\s*(?:[0-9]{1,3}[.)]\s+|\([0-9]{1,3}\)\s*)", "", q.title) if numbered else q.title
            prefix = f"Q{number}. {title}" if numbered else title
            if str(q.id) not in selected:
                continue
            def add(label: str, extract: Callable[[Any], Any], question: Question = q) -> None:
                columns.append(ExportColumn(label, str(question.id), section.title, question.title,
                    lambda answers, question=question, extract=extract: extract(answers.get(str(question.id)))))
            if request.mode == "raw":
                add(prefix, lambda value: _json(value) if isinstance(value, (dict, list)) else (value if value is not None else ""))
            elif request.mode == "analysis" and q.type.value in {"multiple", "multiple_choice"}:
                for option in q.options:
                    add(f"{prefix} [{option.label}]", lambda value, option=option: int(str(option.id) in [str(item) for item in (value if isinstance(value, list) else ([] if value is None else [value]))]))
            elif request.mode == "analysis" and q.type.value in {"matrix", "matrix_single", "matrix_multiple"}:
                for row_index, row in enumerate(q.settings.rows):
                    if q.type.value == "matrix_multiple" or (q.settings.model_extra or {}).get("selectionMode") == "multiple":
                        for column_index, option in enumerate(q.settings.columns):
                            def multiple_cell(value, row=row, row_index=row_index, option=option, column_index=column_index):
                                raw = (value or {}).get(str(row.id), (value or {}).get(str(row_index), []))
                                tokens = [str(item) for item in (raw if isinstance(raw, list) else [raw])]
                                return int(str(option.id) in tokens or str(column_index) in tokens)
                            add(f"{prefix} [{row.label} / {option.label}]", multiple_cell)
                    else:
                        add(f"{prefix} [{row.label}]", lambda value, row=row, row_index=row_index, q=q: _label(q.settings.columns, (value or {}).get(str(row.id), (value or {}).get(str(row_index)))))
            else:
                add(prefix, lambda value, q=q: _display(q, value))
    if len(columns) + 5 > 16384:
        raise DomainValidationError("엑셀 최대 열 수를 초과했습니다. 현재 보기형으로 문항을 나누어 내려받아 주세요.")
    return columns


def _selected_responses(store: SurveyStore, user: UserRecord, survey_id: str, ids: list[str] | None) -> list[SurveyResponseView]:
    if ids == []:
        return []
    wanted = set(ids) if ids is not None else None
    selected: list[SurveyResponseView] = []
    for offset in range(0, 1_048_577, 200):
        page = store.responses(user, survey_id, offset=offset, limit=200)
        selected.extend(item for item in page if wanted is None or item.id in wanted)
        if len(selected) > 1_048_575:
            raise DomainValidationError("엑셀 최대 행 수를 초과했습니다. 조회 조건을 나누어 주세요.")
        if len(page) < 200 or (wanted is not None and len(selected) == len(wanted)):
            break
    if wanted is not None:
        by_id = {item.id: item for item in selected}
        if set(by_id) != wanted:
            raise DomainValidationError("조회 대상 응답이 변경되었습니다. 응답 탐색을 새로고침해 주세요.")
        return [by_id[identifier] for identifier in ids]
    return selected


def _write(sheet, values: list[Any]) -> None:
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    for value in values:
        if isinstance(value, str) and (len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value)):
            raise DomainValidationError("엑셀 셀의 문자 수 또는 문자 형식 제한을 초과했습니다. CSV 원본형을 이용해 주세요.")
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            # Preserve the exact text while preventing formula execution.
            cell.data_type = "s"
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def render_response_workbook(store: SurveyStore, user: UserRecord, survey_id: str, request: ResponseWorkbookRequest) -> tuple[bytes, str]:
    # Both the project and every response page use the regular owner/tenant checks.
    survey = store.get_survey(user, survey_id)
    responses = _selected_responses(store, user, survey_id, request.response_ids)
    columns = _columns(survey, request)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "응답 데이터"
    _write(sheet, ["응답ID", "상태", "제출시각(UTC)", "소요시간(초)", "설문버전"] + [column.label for column in columns])
    for response in responses:
        answers = {str(item.question_id): item.value for item in response.answers}
        _write(sheet, [response.id, "완료", response.submitted_at.astimezone(timezone.utc).isoformat(), "", response.survey_version] + [column.value(answers) for column in columns])
    sheet.freeze_panes = "F2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 58
    for cell in sheet[1]:
        cell.font = Font(name="맑은 고딕", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2457D6")
    for index in range(1, len(columns) + 6):
        sheet.column_dimensions[get_column_letter(index)].width = 34 if index > 5 else [42, 10, 26, 15, 12][index - 1]
    dictionary = workbook.create_sheet("문항 사전")
    _write(dictionary, ["엑셀 열", "안정적인 문항 ID", "섹션", "원문 제목"])
    for column in columns:
        _write(dictionary, [column.label, column.question_id, column.section, column.original_title])
    dictionary.freeze_panes = "A2"
    dictionary.auto_filter.ref = dictionary.dimensions
    for letter, width in {"A": 60, "B": 32, "C": 30, "D": 60}.items():
        dictionary.column_dimensions[letter].width = width
    metadata = workbook.create_sheet("내보내기 정보")
    for values in [["항목", "내용"], ["설문", survey.title], ["프로젝트 ID", survey_id], ["형식", request.mode],
                   ["응답 수", len(responses)], ["문항 열 수", len(columns)], ["시간 기준", "UTC"],
                   ["조회 범위", "응답 탐색의 현재 필터에 해당하는 전체 페이지 응답"],
                   ["값 표기", "미응답은 빈 셀, 복수응답 분석형은 보기별 1/0, 원본형 객체는 JSON"],
                   ["버전", "각 행의 설문버전은 실제 응답을 수집한 확정 버전입니다."]]:
        _write(metadata, values)
    metadata.column_dimensions["A"].width = 22
    metadata.column_dimensions["B"].width = 90
    output = io.BytesIO()
    workbook.save(output)
    stem = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "_", survey.title).strip(" .")[:80] or "survey"
    return output.getvalue(), f"{stem}_responses_{request.mode}.xlsx"
