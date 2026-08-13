from __future__ import annotations

import base64
import html
import io
import logging
import os
import re
import threading
from datetime import date
from functools import lru_cache
from pathlib import Path
from typing import Literal

from backend.errors import ServiceUnavailableError
from backend.models import QuestionReport, SurveyReport


logger = logging.getLogger(__name__)

ExportFormat = Literal["xlsx", "pdf"]
PROJECT_ROOT = Path(__file__).resolve().parent.parent
REPORT_FONT_FAMILY = "IN-SURVEYONE Noto Sans KR"
_FONT_ROOT = PROJECT_ROOT / "node_modules" / "@fontsource-variable" / "noto-sans-kr"
_FONT_CSS = _FONT_ROOT / "index.css"
_FONT_URL_PATTERN = re.compile(r"url\(\./files/([^)]+\.woff2)\)")
_PDF_SEMAPHORE = threading.BoundedSemaphore(2)
_PDF_MAX_HTML_BYTES = 12 * 1024 * 1024
_PDF_TIMEOUT_MS = 30_000


def safe_report_stem(title: str, survey_id: str) -> str:
    normalized = re.sub(r'[\\/:*?"<>|\x00-\x1f]+', "_", title).strip(" ._")
    normalized = re.sub(r"\s+", " ", normalized)[:80]
    fallback = re.sub(r"[^A-Za-z0-9_-]+", "", survey_id)[:24] or "survey"
    return normalized or f"report-{fallback}"


def export_filename(report: SurveyReport, export_format: ExportFormat) -> str:
    stamp = report.generated_at.astimezone().strftime("%Y%m%d-%H%M")
    return f"{safe_report_stem(report.title, report.survey_id)}_{stamp}.{export_format}"


def _safe_excel_value(value: object) -> object:
    if not isinstance(value, str):
        return value
    # Prevent a title, option, or free-text sample from becoming an Excel formula.
    return f"'{value}" if re.match(r"^[\t\r ]*[=+\-@]", value) else value


def _question_type_label(question: QuestionReport) -> str:
    value = question.type.value
    return {
        "single": "객관식",
        "single_choice": "객관식",
        "multiple": "복수응답",
        "multiple_choice": "복수응답",
        "dropdown": "드롭다운",
        "short_text": "단답형",
        "long_text": "장문형",
        "number": "숫자",
        "scale": "척도",
        "slider": "슬라이더",
        "matrix": "매트릭스",
        "table": "표",
        "repeat_table": "반복표",
        "date": "날짜",
    }.get(value, value)


def _primary_result(question: QuestionReport) -> str:
    if question.option_counts:
        label, count = max(question.option_counts.items(), key=lambda item: item[1])
        ratio = (count / question.response_count * 100) if question.response_count else 0
        return f"{label} {count:,}건 ({ratio:.1f}%)"
    if question.average is not None:
        return f"평균 {question.average:,.2f}"
    if question.text_samples:
        return f"주관식 표본 {len(question.text_samples):,}건"
    return "집계 결과 없음"


def render_report_xlsx(
    report: SurveyReport,
    *,
    cutoff_date: date | None,
    response_status: str,
) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as error:  # pragma: no cover - deployment dependency guard
        raise ServiceUnavailableError(
            "Excel 내보내기 라이브러리가 설치되지 않았습니다. 서버 빌드를 확인해 주세요."
        ) from error

    workbook = Workbook()
    workbook.properties.title = report.title
    workbook.properties.subject = "IN-SURVEYONE 기본형 보고서"
    workbook.properties.creator = "IN-SURVEYONE"
    workbook.properties.description = "선택한 집계 기준으로 생성한 기본형 설문 보고서"

    blue = "2457D6"
    dark = "10243E"
    pale = "EDF3FF"
    border_color = "D7E0EE"
    thin = Side(style="thin", color=border_color)

    def style_header(row) -> None:
        for cell in row:
            cell.font = Font(name="맑은 고딕", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=blue)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=thin)

    def auto_width(sheet, minimum: int = 10, maximum: int = 60) -> None:
        for column_cells in sheet.columns:
            letter = get_column_letter(column_cells[0].column)
            longest = max(
                (len(str(cell.value)) for cell in column_cells if cell.value is not None),
                default=minimum,
            )
            sheet.column_dimensions[letter].width = min(max(longest + 3, minimum), maximum)

    summary = workbook.active
    summary.title = "보고서 요약"
    summary.sheet_view.showGridLines = False
    summary.merge_cells("A1:D1")
    summary["A1"] = _safe_excel_value(report.title)
    summary["A1"].font = Font(name="맑은 고딕", size=20, bold=True, color=dark)
    summary["A1"].alignment = Alignment(vertical="center")
    summary.row_dimensions[1].height = 34
    basis_rows = [
        ("보고서 종류", "기본형 보고서"),
        ("설문 프로젝트 ID", report.survey_id),
        ("배포 버전", report.published_version or "미배포 초안"),
        ("응답 기준일", cutoff_date.isoformat() if cutoff_date else "전체 기간"),
        ("응답 상태", {"all": "전체 응답", "complete": "완료 응답", "partial": "진행 중 응답"}[response_status]),
        ("생성 시각", report.generated_at.astimezone().strftime("%Y-%m-%d %H:%M:%S %Z")),
    ]
    for row_index, (label, value) in enumerate(basis_rows, start=3):
        summary.cell(row=row_index, column=1, value=label)
        summary.cell(row=row_index, column=2, value=_safe_excel_value(value))
        summary.cell(row=row_index, column=1).font = Font(name="맑은 고딕", bold=True, color=dark)
        summary.cell(row=row_index, column=1).fill = PatternFill("solid", fgColor=pale)
        summary.cell(row=row_index, column=1).border = Border(bottom=thin)
        summary.cell(row=row_index, column=2).border = Border(bottom=thin)

    metric_row = 11
    metrics = [
        ("전체 응답", report.total_responses),
        ("완료 응답", report.completed_responses),
        ("진행 중 응답", report.partial_responses),
        ("완료율", report.completion_rate),
    ]
    for column, (label, value) in enumerate(metrics, start=1):
        label_cell = summary.cell(row=metric_row, column=column, value=label)
        value_cell = summary.cell(row=metric_row + 1, column=column, value=value)
        label_cell.font = Font(name="맑은 고딕", bold=True, color="FFFFFF")
        label_cell.fill = PatternFill("solid", fgColor=blue)
        label_cell.alignment = Alignment(horizontal="center")
        value_cell.font = Font(name="맑은 고딕", size=16, bold=True, color=dark)
        value_cell.fill = PatternFill("solid", fgColor="F7F9FC")
        value_cell.alignment = Alignment(horizontal="center")
        label_cell.border = value_cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    summary.cell(row=metric_row + 1, column=4).number_format = "0.0%"
    summary.freeze_panes = "A3"
    summary.column_dimensions["A"].width = 23
    summary.column_dimensions["B"].width = 36
    summary.column_dimensions["C"].width = 23
    summary.column_dimensions["D"].width = 23

    questions = workbook.create_sheet("문항별 결과")
    questions.sheet_view.showGridLines = False
    questions.freeze_panes = "A2"
    questions.auto_filter.ref = f"A1:H{max(len(report.questions) + 1, 2)}"
    question_headers = ["번호", "문항 ID", "문항", "유형", "응답", "무응답", "평균", "주요 결과"]
    questions.append(question_headers)
    style_header(questions[1])
    for index, question in enumerate(report.questions, start=1):
        questions.append(
            [
                index,
                _safe_excel_value(question.question_id),
                _safe_excel_value(question.title),
                _question_type_label(question),
                question.response_count,
                max(report.total_responses - question.response_count, 0),
                question.average,
                _safe_excel_value(_primary_result(question)),
            ]
        )
    auto_width(questions)

    distribution = workbook.create_sheet("응답 분포")
    distribution.sheet_view.showGridLines = False
    distribution.freeze_panes = "A2"
    distribution.append(["문항 번호", "문항 ID", "문항", "보기", "응답 수", "문항 내 비율"])
    style_header(distribution[1])
    distribution_row = 2
    for index, question in enumerate(report.questions, start=1):
        for label, count in question.option_counts.items():
            ratio = count / question.response_count if question.response_count else 0
            distribution.append(
                [index, _safe_excel_value(question.question_id), _safe_excel_value(question.title), _safe_excel_value(label), count, ratio]
            )
            distribution.cell(row=distribution_row, column=6).number_format = "0.0%"
            distribution_row += 1
    if distribution_row == 2:
        distribution.append([None, None, "선택형 문항의 응답 분포가 없습니다."])
    distribution.auto_filter.ref = f"A1:F{max(distribution_row - 1, 2)}"
    auto_width(distribution)

    samples = workbook.create_sheet("주관식 표본")
    samples.sheet_view.showGridLines = False
    samples.freeze_panes = "A2"
    samples.append(["문항 번호", "문항 ID", "문항", "표본 번호", "응답 내용"])
    style_header(samples[1])
    sample_row = 2
    for index, question in enumerate(report.questions, start=1):
        for sample_index, sample in enumerate(question.text_samples, start=1):
            samples.append(
                [index, _safe_excel_value(question.question_id), _safe_excel_value(question.title), sample_index, _safe_excel_value(sample)]
            )
            samples.cell(row=sample_row, column=5).alignment = Alignment(wrap_text=True, vertical="top")
            sample_row += 1
    if sample_row == 2:
        samples.append([None, None, "주관식 응답 표본이 없습니다."])
    samples.auto_filter.ref = f"A1:E{max(sample_row - 1, 2)}"
    auto_width(samples, maximum=80)

    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


@lru_cache(maxsize=1)
def embedded_korean_font_css() -> str:
    if not _FONT_CSS.is_file():
        raise ServiceUnavailableError(
            "PDF용 한글 글꼴이 설치되지 않았습니다. bash scripts/render-build.sh로 다시 빌드해 주세요."
        )
    css = _FONT_CSS.read_text(encoding="utf-8")
    encoded: dict[str, str] = {}

    def replace_url(match: re.Match[str]) -> str:
        file_name = match.group(1)
        font_file = _FONT_ROOT / "files" / file_name
        if not font_file.is_file():
            raise ServiceUnavailableError("PDF용 한글 글꼴 파일이 완전하지 않습니다.")
        if file_name not in encoded:
            encoded[file_name] = base64.b64encode(font_file.read_bytes()).decode("ascii")
        return f"url(data:font/woff2;base64,{encoded[file_name]})"

    embedded = _FONT_URL_PATTERN.sub(replace_url, css)
    embedded = embedded.replace("'Noto Sans KR Variable'", f"'{REPORT_FONT_FAMILY}'")
    return embedded.replace("font-display: swap", "font-display: block")


def render_report_html(
    report: SurveyReport,
    *,
    cutoff_date: date | None,
    response_status: str,
) -> str:
    def escape(value: object) -> str:
        return html.escape(str(value), quote=True)

    status_label = {"all": "전체 응답", "complete": "완료 응답", "partial": "진행 중 응답"}[response_status]
    cutoff_label = cutoff_date.isoformat() if cutoff_date else "전체 기간"
    question_sections: list[str] = []
    for index, question in enumerate(report.questions, start=1):
        result_rows = []
        for label, count in question.option_counts.items():
            ratio = (count / question.response_count * 100) if question.response_count else 0
            result_rows.append(
                '<div class="bar-row">'
                f'<div class="bar-label"><span>{escape(label)}</span><strong>{count:,}건 · {ratio:.1f}%</strong></div>'
                f'<div class="bar-track"><span style="width:{min(max(ratio, 0), 100):.2f}%"></span></div>'
                "</div>"
            )
        if question.average is not None:
            result_rows.append(f'<p class="numeric-result">평균 <strong>{question.average:,.2f}</strong></p>')
        if question.text_samples:
            result_rows.append(
                '<ol class="samples">'
                + "".join(f"<li>{escape(sample)}</li>" for sample in question.text_samples)
                + "</ol>"
            )
        if not result_rows:
            result_rows.append('<p class="empty">집계할 응답이 없습니다.</p>')
        missing = max(report.total_responses - question.response_count, 0)
        question_sections.append(
            '<section class="question">'
            f'<div class="question-index">Q{index:02d}</div>'
            '<div class="question-body">'
            f'<h2>{escape(question.title)}</h2>'
            f'<p class="question-meta">{escape(_question_type_label(question))} · 응답 {question.response_count:,}건 · 무응답 {missing:,}건</p>'
            + "".join(result_rows)
            + "</div></section>"
        )

    font_css = embedded_korean_font_css()
    version_label = f"v{report.published_version}" if report.published_version else "미배포 초안"
    completion_percent = report.completion_rate * 100
    return f"""<!doctype html>
<html lang="ko">
<head>
<meta charset="utf-8">
<title>{escape(report.title)} - 기본형 보고서</title>
<style>
{font_css}
@page {{ size: A4; margin: 16mm 14mm 17mm; }}
* {{ box-sizing: border-box; -webkit-print-color-adjust: exact; print-color-adjust: exact; }}
body {{ margin: 0; color: #10243e; font-family: '{REPORT_FONT_FAMILY}', sans-serif; font-size: 10.5pt; line-height: 1.55; }}
.cover {{ min-height: 250mm; display: flex; flex-direction: column; justify-content: space-between; page-break-after: always; }}
.brand {{ color: #2457d6; font-size: 10pt; font-weight: 800; letter-spacing: .12em; }}
.cover h1 {{ margin: 32mm 0 5mm; font-size: 28pt; line-height: 1.25; word-break: keep-all; }}
.cover-lead {{ color: #53657b; font-size: 12pt; }}
.basis {{ margin-top: 20mm; padding: 7mm; background: #edf3ff; border: 1px solid #ccdaf3; border-radius: 4mm; }}
.basis dl {{ display: grid; grid-template-columns: 34mm 1fr; gap: 2.5mm 4mm; margin: 0; }}
.basis dt {{ color: #5e7088; }} .basis dd {{ margin: 0; font-weight: 700; }}
.cover-footer {{ display: flex; justify-content: space-between; color: #66778d; border-top: 1px solid #d8e1ef; padding-top: 4mm; }}
.section-title {{ margin: 0 0 5mm; font-size: 18pt; }}
.kpis {{ display: grid; grid-template-columns: repeat(4, 1fr); gap: 3mm; margin-bottom: 8mm; }}
.kpi {{ padding: 4mm; border: 1px solid #d8e1ef; border-radius: 3mm; background: #f7f9fc; }}
.kpi span {{ display: block; color: #63748b; font-size: 8.5pt; }} .kpi strong {{ display: block; margin-top: 2mm; font-size: 17pt; }}
.question {{ display: grid; grid-template-columns: 14mm 1fr; gap: 4mm; margin-bottom: 5mm; padding: 5mm; border: 1px solid #d8e1ef; border-radius: 3mm; break-inside: avoid; }}
.question-index {{ color: #2457d6; font-weight: 800; }} .question h2 {{ margin: 0; font-size: 13pt; }}
.question-meta {{ margin: 1mm 0 4mm; color: #697a90; font-size: 8.5pt; }}
.bar-row {{ margin: 2.5mm 0; }} .bar-label {{ display: flex; justify-content: space-between; gap: 5mm; margin-bottom: 1mm; font-size: 9pt; }}
.bar-label strong {{ white-space: nowrap; }} .bar-track {{ height: 3mm; overflow: hidden; border-radius: 2mm; background: #e7edf6; }}
.bar-track span {{ display: block; height: 100%; min-width: .5mm; background: #2457d6; border-radius: inherit; }}
.numeric-result {{ padding: 4mm; background: #edf3ff; border-radius: 2.5mm; }} .numeric-result strong {{ color: #2457d6; font-size: 16pt; }}
.samples {{ margin: 0; padding-left: 6mm; }} .samples li {{ margin-bottom: 2mm; padding: 2.5mm; background: #f7f9fc; border-radius: 2mm; word-break: break-word; }}
.empty {{ color: #77869a; }}
</style>
</head>
<body>
<section class="cover">
  <div>
    <div class="brand">IN-SURVEYONE · BASIC REPORT</div>
    <h1>{escape(report.title)}</h1>
    <p class="cover-lead">설문 응답의 핵심 지표와 문항별 분포를 선택한 기준으로 집계한 보고서입니다.</p>
    <div class="basis"><dl>
      <dt>프로젝트 ID</dt><dd>{escape(report.survey_id)}</dd>
      <dt>배포 버전</dt><dd>{escape(version_label)}</dd>
      <dt>응답 기준일</dt><dd>{escape(cutoff_label)}</dd>
      <dt>응답 상태</dt><dd>{escape(status_label)}</dd>
      <dt>생성 시각</dt><dd>{escape(report.generated_at.astimezone().strftime('%Y-%m-%d %H:%M:%S %Z'))}</dd>
    </dl></div>
  </div>
  <div class="cover-footer"><span>IN-SURVEYONE</span><span>기본형 보고서</span></div>
</section>
<main>
  <h2 class="section-title">응답 핵심 지표</h2>
  <div class="kpis">
    <div class="kpi"><span>전체 응답</span><strong>{report.total_responses:,}</strong></div>
    <div class="kpi"><span>완료 응답</span><strong>{report.completed_responses:,}</strong></div>
    <div class="kpi"><span>진행 중 응답</span><strong>{report.partial_responses:,}</strong></div>
    <div class="kpi"><span>완료율</span><strong>{completion_percent:.1f}%</strong></div>
  </div>
  <h2 class="section-title">문항별 응답 결과</h2>
  {''.join(question_sections) if question_sections else '<p class="empty">표시할 문항 집계가 없습니다.</p>'}
</main>
</body>
</html>"""


def render_report_pdf(html_document: str) -> bytes:
    encoded = html_document.encode("utf-8")
    if len(encoded) > _PDF_MAX_HTML_BYTES:
        raise ServiceUnavailableError("PDF 보고서가 안전하게 생성할 수 있는 크기를 초과했습니다.")
    if not _PDF_SEMAPHORE.acquire(timeout=10):
        raise ServiceUnavailableError("PDF 생성 요청이 많습니다. 잠시 후 다시 시도해 주세요.")
    try:
        os.environ.setdefault("PLAYWRIGHT_BROWSERS_PATH", "0")
        try:
            from playwright.sync_api import Error as PlaywrightError
            from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
            from playwright.sync_api import sync_playwright
        except ImportError as error:  # pragma: no cover - deployment dependency guard
            raise ServiceUnavailableError(
                "PDF 내보내기 라이브러리가 설치되지 않았습니다. 서버 빌드를 확인해 주세요."
            ) from error

        browser = None
        try:
            with sync_playwright() as playwright:
                browser = playwright.chromium.launch(
                    headless=True,
                    args=["--disable-dev-shm-usage"],
                )
                context = browser.new_context(java_script_enabled=False, service_workers="block")
                page = context.new_page()

                def block_external(route) -> None:
                    url = route.request.url
                    if url == "about:blank" or url.startswith("data:"):
                        route.continue_()
                    else:
                        route.abort("blockedbyclient")

                page.route("**/*", block_external)
                page.set_content(html_document, wait_until="load", timeout=_PDF_TIMEOUT_MS)
                page.evaluate("async () => { await document.fonts.ready; }")
                font_ready = page.evaluate(
                    f"() => document.fonts.check('16px \\\"{REPORT_FONT_FAMILY}\\\"', '한글 보고서')"
                )
                if not font_ready:
                    raise ServiceUnavailableError("PDF용 한글 글꼴을 불러오지 못했습니다.")
                pdf = page.pdf(
                    format="A4",
                    print_background=True,
                    prefer_css_page_size=True,
                    margin={"top": "0", "right": "0", "bottom": "0", "left": "0"},
                )
                context.close()
                browser.close()
                browser = None
                if not pdf.startswith(b"%PDF-"):
                    raise ServiceUnavailableError("PDF 파일 형식 검증에 실패했습니다.")
                return pdf
        except ServiceUnavailableError:
            raise
        except PlaywrightTimeoutError as error:
            logger.warning("M05001 PDF rendering timed out: %s", error)
            raise ServiceUnavailableError("PDF 생성 시간이 초과되었습니다. 잠시 후 다시 시도해 주세요.") from error
        except PlaywrightError as error:
            logger.error("M05001 PDF rendering failed: %s", error)
            raise ServiceUnavailableError(
                "PDF 생성기를 실행하지 못했습니다. bash scripts/render-build.sh 빌드를 확인해 주세요."
            ) from error
        finally:
            if browser is not None:
                try:
                    browser.close()
                except Exception:  # pragma: no cover - best-effort process cleanup
                    logger.warning("M05001 PDF browser cleanup failed", exc_info=True)
    finally:
        _PDF_SEMAPHORE.release()
