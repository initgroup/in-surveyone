from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import ValidationError

from backend.errors import DomainValidationError
from backend.models import DeliveryRecipientInput


MAX_RECIPIENT_FILE_BYTES = 4 * 1024 * 1024
MAX_RECIPIENT_ROWS = 2_000
RECIPIENT_HEADERS = (
    "name",
    "countryCode",
    "preferredLanguage",
    "email",
    "mobile",
    "fax",
    "kakaoId",
    "channels",
    "audienceCode",
)
CONTACT_CHANNELS = {
    "email": "EMAIL",
    "mobile": "SMS",
    "kakaoId": "KAKAO",
    "fax": "FAX",
}


def _decode_text(data: bytes) -> tuple[str, str]:
    signatures = (
        (b"\xef\xbb\xbf", "utf-8-sig"),
        (b"\xff\xfe", "utf-16"),
        (b"\xfe\xff", "utf-16"),
    )
    for signature, encoding in signatures:
        if data.startswith(signature):
            return data.decode(encoding), encoding
    if b"\x00" in data[:200]:
        try:
            return data.decode("utf-16"), "utf-16"
        except UnicodeDecodeError:
            pass
    for encoding in ("utf-8", "cp949", "euc-kr", "cp932", "gb18030", "big5"):
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise DomainValidationError(
        "텍스트 파일 인코딩을 확인할 수 없습니다. UTF-8, UTF-16 또는 한글 ANSI(CP949)로 저장해 주세요."
    )


def _text_rows(data: bytes) -> tuple[list[list[Any]], str]:
    text, encoding = _decode_text(data)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel
    rows = [list(row) for row in csv.reader(io.StringIO(text), dialect)]
    return rows, encoding


def _xlsx_rows(data: bytes) -> list[list[Any]]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except Exception as exc:
        raise DomainValidationError("XLSX 파일을 읽을 수 없습니다.") from exc
    try:
        sheet = workbook.active
        rows: list[list[Any]] = []
        for row in sheet.iter_rows():
            if any(cell.data_type == "f" for cell in row):
                raise DomainValidationError("대상자 XLSX에는 수식을 사용할 수 없습니다.")
            rows.append([cell.value for cell in row])
            if len(rows) > MAX_RECIPIENT_ROWS + 1:
                raise DomainValidationError(
                    f"한 번에 가져올 수 있는 대상자는 최대 {MAX_RECIPIENT_ROWS:,}명입니다."
                )
        return rows
    finally:
        workbook.close()


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_recipient_file(filename: str, data: bytes) -> tuple[list[DeliveryRecipientInput], str, str | None]:
    if not data:
        raise DomainValidationError("가져올 파일이 비어 있습니다.")
    if len(data) > MAX_RECIPIENT_FILE_BYTES:
        raise DomainValidationError("대상자 파일은 4MB 이하여야 합니다.")
    extension = Path(filename or "").suffix.lower()
    if extension == ".xlsx":
        rows = _xlsx_rows(data)
        file_type, encoding = "xlsx", None
    elif extension in {".csv", ".txt"}:
        rows, encoding = _text_rows(data)
        file_type = extension.removeprefix(".")
    else:
        raise DomainValidationError("CSV, TXT 또는 XLSX 파일만 가져올 수 있습니다.")
    rows = [row for row in rows if any(_cell(value) for value in row)]
    if len(rows) < 2:
        raise DomainValidationError("파일에 대상자 데이터가 없습니다.")

    raw_headers = [_cell(value).lstrip("\ufeff") for value in rows[0]]
    header_lookup = {value.casefold(): value for value in RECIPIENT_HEADERS}
    header_lookup.update({"이름": "name", "대상자명": "name", "휴대전화": "mobile", "휴대폰": "mobile", "전화번호": "mobile", "이메일": "email", "국가": "countryCode", "언어": "preferredLanguage", "대상구분": "audienceCode"})
    headers = [header_lookup.get(value.casefold(), value) for value in raw_headers]
    if len(headers) != len(set(headers)):
        raise DomainValidationError("파일에 같은 열 이름이 중복되어 있습니다.")
    required = {"name", "countryCode", "preferredLanguage"}
    missing = sorted(required - set(headers))
    if missing:
        raise DomainValidationError("필수 열이 없습니다: " + ", ".join(missing))
    if len(rows) - 1 > MAX_RECIPIENT_ROWS:
        raise DomainValidationError(
            f"한 번에 가져올 수 있는 대상자는 최대 {MAX_RECIPIENT_ROWS:,}명입니다."
        )

    recipients: list[DeliveryRecipientInput] = []
    errors: list[str] = []
    for row_number, row in enumerate(rows[1:], start=2):
        values = {
            header: _cell(row[index]) if index < len(row) else ""
            for index, header in enumerate(headers)
            if header in RECIPIENT_HEADERS
        }
        channels = [
            value.strip().upper()
            for value in values.get("channels", "").replace(",", "|").replace(";", "|").split("|")
            if value.strip()
        ]
        if values.get("mobile") and not values.get("kakaoId"):
            values["kakaoId"] = values["mobile"]
        if not channels:
            channels = [
                channel for field, channel in CONTACT_CHANNELS.items() if values.get(field)
            ]
        try:
            recipients.append(
                DeliveryRecipientInput.model_validate(
                    {
                        **values,
                        "audienceCode": values.get("audienceCode") or None,
                        "channels": channels,
                    }
                )
            )
        except ValidationError as exc:
            message = exc.errors()[0].get("msg", "입력값을 확인해 주세요.")
            errors.append(f"{row_number}행: {message}")
            if len(errors) >= 10:
                break
    if errors:
        raise DomainValidationError("대상자 파일을 확인해 주세요. " + " / ".join(errors))
    if not recipients:
        raise DomainValidationError("가져올 대상자 행이 없습니다.")
    return recipients, file_type, encoding


def recipient_xlsx_sample() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "recipients"
    sheet.append(list(RECIPIENT_HEADERS))
    sheet.append(
        [
            "응답자 A001",
            "KR",
            "ko",
            "sample@example.com",
            "010-0000-0000",
            "",
            "",
            "EMAIL|SMS|KAKAO",
            "",
        ]
    )
    fill = PatternFill("solid", fgColor="2457D6")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
    widths = (22, 15, 20, 28, 20, 18, 22, 24, 22)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
