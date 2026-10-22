from __future__ import annotations

from contextlib import contextmanager
from datetime import datetime, timezone
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError

from backend.answer_codes import coded_value
from backend.models import PublicSurvey, Question, SurveySection, SurveySettings, SurveyTheme
from backend.oracle_store import OracleStore
from main import app


def test_choice_and_axis_result_codes_reject_empty_duplicate_and_control_values():
    for value in ["", " ", "bad\ncode", "a" * 121]:
        with pytest.raises(ValidationError):
            Question.model_validate({"type": "single", "options": [{"id": "a", "label": "A", "resultCode": value}]})
    for field in ["options", "rows", "columns"]:
        items = [{"id": "a", "label": "A", "resultCode": "same"}, {"id": "b", "label": "B", "resultCode": "same"}]
        payload = {"type": "single", "options": items} if field == "options" else {"type": "table", "settings": {field: items}}
        with pytest.raises(ValidationError):
            Question.model_validate(payload)
    with pytest.raises(ValidationError):
        Question.model_validate({"type": "scale", "min": 1, "max": 3, "settings": {"scalePoints": [{"score": 1, "label": "A", "resultCode": "x"}, {"score": 2, "label": "B", "resultCode": "x"}]}})
    with pytest.raises(ValidationError):
        Question.model_validate({"type": "scale", "min": 1, "max": 3, "settings": {"scalePoints": [{"score": 4, "label": "A", "resultCode": "x"}]}})


@pytest.mark.parametrize("kind", ["matrix_single", "matrix_multiple", "table", "repeat_table"])
def test_matrix_and_table_codes_keep_zero_and_distinguish_labels_from_codes(kind):
    question = Question.model_validate({"type": kind, "settings": {"rows": [{"id": "r", "label": "Row label", "resultCode": "R10"}], "columns": [{"id": "c", "label": "Column label", "resultCode": "C20"}]}})
    raw = {"r": ["c"]} if kind == "matrix_multiple" else {"r": "c"} if kind == "matrix_single" else [[0]]
    expected = {"R10": ["C20"]} if kind == "matrix_multiple" else {"R10": "C20"} if kind == "matrix_single" else [{"C20": 0}] if kind == "repeat_table" else {"R10": {"C20": 0}}
    assert coded_value(question, raw) == expected


def test_direct_input_codes_preserve_actual_values_and_legacy_ids():
    for kind, value in [("short_text", "문자 응답"), ("long_text", "긴 응답"), ("number", 0), ("slider", 0), ("date", "2026-10-22")]:
        assert coded_value(Question(type=kind), value) == value
    legacy = Question(type="single", options=[{"id": "stable-a", "label": "Visible"}])
    assert coded_value(legacy, "stable-a") == "stable-a"


def _snapshot(version, code):
    return PublicSurvey(survey_id="survey", slug="synthetic", version=version, title="Synthetic", category="test", sections=[SurveySection(id="s", title="S", questions=[Question(type="single", id="q", options=[{"id": "a", "label": "Visible", "resultCode": code}])])], settings=SurveySettings(), theme=SurveyTheme(), published_at=datetime(2026, 10, 22, tzinfo=timezone.utc))


def test_oracle_responses_resolve_each_version_with_tenant_scoped_binds():
    class Pool:
        @contextmanager
        def connection(self):
            yield object()
    class Crud:
        versions = []
        def fetch_all(self, connection, sql_id, binds):
            assert binds["tenantId"] == "tenant" and binds["ownerUserId"] == "owner"
            if sql_id == "RESPONSE_LIST_BY_OWNER":
                return [{"response_id": f"resp-{v}", "survey_id": "survey", "survey_version_no": v, "submitted_at": datetime(2026, 10, 22, tzinfo=timezone.utc)} for v in [1, 2]]
            assert sql_id == "RESPONSE_ANSWER_LIST_BY_OWNER"
            return [{"response_id": f"resp-{v}", "question_id": "q", "answer_value_json": '"a"'} for v in [1, 2]]
        def fetch_one(self, connection, sql_id, binds):
            assert sql_id == "SURVEY_VERSION_BY_NO"
            assert binds["tenantId"] == "tenant" and binds["surveyId"] == "survey"
            self.versions.append(binds["versionNo"])
            return {"snapshot_json": _snapshot(binds["versionNo"], f"CODE-{binds['versionNo']}").model_dump(mode="json", by_alias=True)}
    class Store(OracleStore):
        def _survey_row(self, connection, user, survey_id):
            assert user.id == "owner" and survey_id == "survey"
            return {}
    store = object.__new__(Store)
    store.pool, store.crud = Pool(), Crud()
    rows = store.responses(SimpleNamespace(id="owner", tenant_id="tenant"), "survey")
    assert [row.coded_answers for row in rows] == [{"q": "CODE-1"}, {"q": "CODE-2"}]
    assert [row.answers[0].value for row in rows] == ["a", "a"]
    assert sorted(store.crud.versions) == [1, 2]


def test_result_codes_roundtrip_submit_republish_and_export_using_response_version():
    with TestClient(app, client=("result-code-audit", 50000)) as client:
        login = "code-" + uuid4().hex[:10]
        signup = client.post("/api/auth/signup", json={"loginId": login, "password": "Synthetic!234", "name": "Test", "email": f"{login}@example.test", "organizationName": "Synthetic", "termsAccepted": True, "privacyAccepted": True})
        assert signup.status_code == 201, signup.text
        created = client.post("/api/surveys", json={"title": "Synthetic result codes", "questions": [{"id": "q", "type": "single", "title": "Question", "options": [{"id": "a", "label": "Visible text", "resultCode": "=CODE-1"}, {"id": "b", "label": "Other", "resultCode": "OTHER"}]}, {"id": "scale", "type": "scale", "title": "Scale", "min": 1, "max": 3, "settings": {"scalePoints": [{"score": 1, "label": "Low", "resultCode": "LOW"}, {"score": 2, "label": "Middle", "resultCode": "MID"}, {"score": 3, "label": "High", "resultCode": "HIGH"}]}}]})
        assert created.status_code == 201, created.text
        survey_id = created.json()["data"]["id"]
        published = client.post(f"/api/surveys/{survey_id}/publish", json={})
        assert published.status_code == 200, published.text
        slug = published.json()["data"]["slug"]
        def submit():
            result = client.post(f"/api/public/surveys/{slug}/responses", json={"answers": {"q": "a", "scale": 2}})
            assert result.status_code == 201, result.text
        submit()
        draft = client.get(f"/api/surveys/{survey_id}").json()["data"]
        draft["sections"][0]["questions"][0]["options"][0]["resultCode"] = "CODE-2"
        draft["sections"][0]["questions"][1]["settings"]["scalePoints"][1]["resultCode"] = "MIDDLE-2"
        updated = client.put(f"/api/surveys/{survey_id}", json={"expectedRevision": draft["revision"], "sections": draft["sections"]})
        assert updated.status_code == 200, updated.text
        republished = client.post(f"/api/surveys/{survey_id}/publish", json={})
        assert republished.status_code == 200, republished.text
        submit()
        rows = client.get(f"/api/surveys/{survey_id}/responses").json()["data"]
        by_version = {row["surveyVersion"]: row for row in rows}
        assert by_version[1]["codedAnswers"] == {"q": "=CODE-1", "scale": "MID"}
        assert by_version[2]["codedAnswers"] == {"q": "CODE-2", "scale": "MIDDLE-2"}
        assert all(next(answer for answer in row["answers"] if answer["questionId"] == "q")["value"] == "a" for row in rows)
        exported = client.post(f"/api/surveys/{survey_id}/responses/export", json={"mode": "codes"})
        assert exported.status_code == 200, exported.text
        sheet = load_workbook(BytesIO(exported.content)).active
        assert {sheet.cell(row, 6).value for row in [2, 3]} == {"=CODE-1", "CODE-2"}
        assert all(sheet.cell(row, 6).data_type == "s" for row in [2, 3])
