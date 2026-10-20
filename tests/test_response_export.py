from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.errors import DomainValidationError
from backend.models import AnswerItem, Question, ResponseWorkbookRequest, SurveyResponseView, SurveySection, SurveySettings
from backend.response_export import render_response_workbook
from main import app


def _fixture():
    questions = [
        Question(id="zero", type="number", title="Zero"),
        Question(id="text", type="short_text", title="Formula text"),
        Question(id="multi", type="multiple", title="Multiple", options=[{"id": "a", "label": "Alpha"}, {"id": "b", "label": "Beta"}]),
        Question(id="matrix", type="matrix_multiple", title="Matrix", settings={"rows": [{"id": "r", "label": "Row"}], "columns": [{"id": "c", "label": "Column"}, {"id": "d", "label": "Other"}]}),
        Question(id="table", type="table", title="Table", settings={"rows": [{"id": "r", "label": "Branch"}], "columns": [{"id": "c", "label": "Count"}]}),
    ]
    survey = SimpleNamespace(title="Synthetic survey", sections=[SurveySection(title="Section", questions=questions)], settings=SurveySettings())
    records = [SurveyResponseView(id=f"response-{index}", survey_id="survey-test", survey_version=2, metadata={}, submitted_at=datetime(2026, 10, 20, tzinfo=timezone.utc), answers=[
        AnswerItem(question_id="zero", value=0), AnswerItem(question_id="text", value="=1+1"),
        AnswerItem(question_id="multi", value=["a"]), AnswerItem(question_id="matrix", value={"r": ["c"]}),
        AnswerItem(question_id="table", value=[[0]]),
    ]) for index in range(250)]
    class Store:
        offsets = []
        def get_survey(self, actor, identifier):
            assert actor == "owner" and identifier == "survey-test"
            return survey
        def responses(self, actor, identifier, *, offset, limit):
            assert actor == "owner" and identifier == "survey-test"
            self.offsets.append(offset)
            return records[offset:offset + limit]
    return Store(), records


def test_workbook_preserves_all_pages_types_and_semantic_columns():
    store, records = _fixture()
    content, name = render_response_workbook(store, "owner", "survey-test", ResponseWorkbookRequest(mode="visible"))
    workbook = load_workbook(BytesIO(content))
    sheet = workbook["응답 데이터"]
    assert sheet.max_row == 251 and sheet.max_column == 10
    assert store.offsets == [0, 200]
    assert sheet["F2"].value == 0 and sheet["F2"].data_type == "n"
    assert sheet["G2"].value == "=1+1" and sheet["G2"].data_type == "s"
    assert sheet["H2"].value == "Alpha"
    assert sheet["I2"].value == "Row: Column"
    assert sheet["J2"].value == "Branch: Count=0"
    assert sheet.freeze_panes == "F2" and sheet.auto_filter.ref == "A1:J251"
    assert workbook.sheetnames == ["응답 데이터", "문항 사전", "내보내기 정보"]
    assert name.endswith(".xlsx")
    for mode, expected in [("analysis", [0, "=1+1", 1, 0, 1, 0, "Branch: Count=0"]), ("raw", [0, "=1+1", '["a"]', '{"r":["c"]}', '[[0]]'])]:
        content, _ = render_response_workbook(store, "owner", "survey-test", ResponseWorkbookRequest(mode=mode, response_ids=[records[-1].id]))
        values = list(load_workbook(BytesIO(content)).active.values)
        assert len(values) == 2
        assert list(values[1][5:]) == expected


def test_workbook_respects_explicit_selection_empty_results_and_rejects_unknown_ids():
    store, records = _fixture()
    request = ResponseWorkbookRequest(question_ids=["zero"], response_ids=[records[-1].id, records[0].id])
    content, _ = render_response_workbook(store, "owner", "survey-test", request)
    rows = list(load_workbook(BytesIO(content)).active.values)
    assert len(rows) == 3 and len(rows[0]) == 6
    assert [row[0] for row in rows[1:]] == [records[-1].id, records[0].id]
    content, _ = render_response_workbook(store, "owner", "survey-test", ResponseWorkbookRequest(response_ids=[], question_ids=[]))
    sheet = load_workbook(BytesIO(content)).active
    assert sheet.max_row == 1 and sheet.max_column == 5
    for request in [ResponseWorkbookRequest(response_ids=["foreign-response"]), ResponseWorkbookRequest(question_ids=["foreign-question"])]:
        with pytest.raises(DomainValidationError):
            render_response_workbook(store, "owner", "survey-test", request)


def _signup(client, prefix):
    identifier = f"{prefix}-{uuid4().hex[:8]}"
    result = client.post("/api/auth/signup", json={"loginId": identifier, "password": "Synthetic!234", "name": "Test", "email": f"{identifier}@example.test", "organizationName": "Synthetic", "termsAccepted": True, "privacyAccepted": True})
    assert result.status_code == 201, result.text


def test_workbook_endpoint_requires_session_project_ownership_and_valid_selection():
    with TestClient(app, client=("response-export-owner", 50000)) as owner:
        _signup(owner, "xlsx-owner")
        created = owner.post("/api/surveys", json={"title": "Synthetic 45 questions", "questions": [{"id": f"q-{index}", "type": "number", "title": f"Question {index}"} for index in range(45)]})
        assert created.status_code == 201, created.text
        survey_id = created.json()["data"]["id"]
        published = owner.post(f"/api/surveys/{survey_id}/publish", json={})
        assert published.status_code == 200, published.text
        slug = published.json()["data"]["slug"]
        submitted = owner.post(f"/api/public/surveys/{slug}/responses", json={"answers": {"q-0": 0, "q-44": 9}})
        assert submitted.status_code == 201, submitted.text
        response_id = submitted.json()["data"]["responseId"]
        endpoint = f"/api/surveys/{survey_id}/responses/export"
        result = owner.post(endpoint, json={"mode": "visible", "responseIds": [response_id]})
        assert result.status_code == 200, result.text
        assert "spreadsheetml" in result.headers["content-type"]
        assert "no-store" in result.headers["cache-control"]
        assert "filename*=UTF-8''" in result.headers["content-disposition"]
        sheet = load_workbook(BytesIO(result.content)).active
        assert sheet.max_column == 50 and sheet.max_row == 2
        assert sheet["F2"].value == 0 and sheet["AX2"].value == 9
        assert owner.post(endpoint, json={"responseIds": ["foreign-response"]}).status_code == 422
        assert owner.post(endpoint, json={"responseIds": [response_id, response_id]}).status_code == 422
        with TestClient(app, client=("response-export-outsider", 50001)) as outsider:
            assert outsider.post(endpoint, json={}).status_code == 401
            _signup(outsider, "xlsx-other")
            assert outsider.post(endpoint, json={}).status_code == 404
