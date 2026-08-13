from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.models import QuestionReport, QuestionType, SurveyReport
from backend.report_export import render_report_html, render_report_xlsx
from main import app


def _signup(client: TestClient, prefix: str) -> None:
    suffix = uuid4().hex[:10]
    response = client.post(
        "/api/auth/signup",
        json={
            "loginId": f"{prefix}-{suffix}",
            "password": "SurveyOne!234",
            "name": "보고서 담당자",
            "email": f"{prefix}-{suffix}@example.test",
            "organizationName": "보고서 시험 기관",
            "termsAccepted": True,
            "privacyAccepted": True,
        },
    )
    assert response.status_code == 201, response.text


def _sample_report() -> SurveyReport:
    return SurveyReport(
        survey_id="survey-report-test",
        title="=위험한 수식 제목",
        total_responses=4,
        completed_responses=3,
        partial_responses=1,
        completion_rate=0.75,
        published_version=2,
        generated_at=datetime(2026, 8, 14, 3, 30, tzinfo=timezone.utc),
        questions=[
            QuestionReport(
                question_id="q-choice",
                title="가장 만족한 기능",
                type=QuestionType.SINGLE_CHOICE,
                response_count=4,
                option_counts={"설문 편집": 3, "=HYPERLINK(\"bad\")": 1},
            ),
            QuestionReport(
                question_id="q-score",
                title="만족도 점수",
                type=QuestionType.SCALE,
                response_count=3,
                average=4.25,
            ),
            QuestionReport(
                question_id="q-text",
                title="개선 의견",
                type=QuestionType.LONG_TEXT,
                response_count=1,
                text_samples=["@악성수식이 아닌 일반 텍스트로 보관"],
            ),
        ],
    )


def test_xlsx_contains_report_sheets_and_escapes_formula_like_text() -> None:
    content = render_report_xlsx(_sample_report(), cutoff_date=None, response_status="all")
    assert content.startswith(b"PK")

    workbook = load_workbook(BytesIO(content), data_only=False)
    assert workbook.sheetnames == ["보고서 요약", "문항별 결과", "응답 분포", "주관식 표본"]
    assert workbook["보고서 요약"]["A1"].value == "'=위험한 수식 제목"
    distribution_values = [cell.value for cell in workbook["응답 분포"]["D"]]
    assert "'=HYPERLINK(\"bad\")" in distribution_values
    sample_values = [cell.value for cell in workbook["주관식 표본"]["E"]]
    assert "'@악성수식이 아닌 일반 텍스트로 보관" in sample_values
    assert workbook["보고서 요약"]["D12"].number_format == "0.0%"
    workbook.close()


def test_pdf_html_is_self_contained_and_escapes_report_values(monkeypatch) -> None:
    monkeypatch.setattr(
        "backend.report_export.embedded_korean_font_css",
        lambda: "@font-face { font-family: 'IN-SURVEYONE Noto Sans KR'; }",
    )
    document = render_report_html(_sample_report(), cutoff_date=None, response_status="complete")
    assert "IN-SURVEYONE · BASIC REPORT" in document
    assert "완료 응답" in document
    assert "설문 편집" in document
    assert "=HYPERLINK(&quot;bad&quot;)" in document
    assert "<script" not in document.lower()
    assert "http://" not in document and "https://" not in document


def test_authenticated_export_endpoint_returns_xlsx_and_pdf(monkeypatch) -> None:
    with TestClient(app) as owner:
        _signup(owner, "report-export-owner")
        created = owner.post(
            "/api/surveys",
            json={"title": "실제 내보내기 보고서", "category": "사회조사"},
        )
        assert created.status_code == 201, created.text
        survey_id = created.json()["data"]["id"]

        xlsx = owner.get(
            f"/api/surveys/{survey_id}/report/export",
            params={"format": "xlsx", "responseStatus": "ALL"},
        )
        assert xlsx.status_code == 200, xlsx.text
        assert xlsx.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert "filename*=UTF-8''" in xlsx.headers["content-disposition"]
        assert xlsx.headers["cache-control"] == "no-store, no-cache, must-revalidate"
        assert xlsx.content.startswith(b"PK")

        monkeypatch.setattr(
            "backend.routers.surveys.render_report_html",
            lambda *args, **kwargs: "<!doctype html><html><body>보고서</body></html>",
        )
        monkeypatch.setattr(
            "backend.routers.surveys.render_report_pdf",
            lambda document: b"%PDF-1.7\nmock-report",
        )
        pdf = owner.get(
            f"/api/surveys/{survey_id}/report/export",
            params={"format": "pdf"},
        )
        assert pdf.status_code == 200, pdf.text
        assert pdf.headers["content-type"].startswith("application/pdf")
        assert pdf.content.startswith(b"%PDF-")

        invalid = owner.get(
            f"/api/surveys/{survey_id}/report/export",
            params={"format": "csv"},
        )
        assert invalid.status_code == 422

        with TestClient(app) as outsider:
            _signup(outsider, "report-export-outsider")
            denied = outsider.get(
                f"/api/surveys/{survey_id}/report/export",
                params={"format": "xlsx"},
            )
            assert denied.status_code == 404


def test_m05001_uses_binary_download_endpoint_and_render_build_contract() -> None:
    project_root = Path(__file__).resolve().parents[1]
    page_script = (project_root / "frontend" / "js" / "M05001.js").read_text(encoding="utf-8")
    page_html = (project_root / "frontend" / "pages" / "M05001.html").read_text(encoding="utf-8")
    build_script = (project_root / "scripts" / "render-build.sh").read_text(encoding="utf-8")

    assert "/report/export?${params}" in page_script
    assert "response.blob()" in page_script
    assert "URL.revokeObjectURL" in page_script
    assert "Beta/Demo 기능" not in page_script
    assert "Excel 집계표 또는 인쇄용 PDF" in page_html
    assert "python -m playwright install --only-shell chromium" in build_script
    assert "python scripts/verify-report-export.py" in build_script
